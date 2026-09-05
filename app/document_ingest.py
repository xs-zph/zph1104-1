"""知识库文档上传与文本提取。

上传文件只在内存中处理，不落盘；提取后的文本交给 RAG 文档集合持久化。
"""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

from app import config, rag

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt"}
MAX_PDF_PAGES = 100


def _clean_text(text: str) -> str:
    text = (text or "").replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError("文档没有可索引的文本内容")
    if len(text) > config.Config.MAX_DOCUMENT_CHARS:
        raise ValueError(f"文档文本不能超过 {config.Config.MAX_DOCUMENT_CHARS} 个字符")
    return text


def _validate(filename: str, data: bytes) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("仅支持 PDF、DOCX、Markdown 或 TXT 文件")
    if not data:
        raise ValueError("上传文件不能为空")
    if len(data) > config.Config.MAX_DOCUMENT_BYTES:
        raise ValueError(f"文件不能超过 {config.Config.MAX_DOCUMENT_BYTES // 1024 // 1024} MB")
    if suffix == ".pdf" and not data.startswith(b"%PDF-"):
        raise ValueError("PDF 文件结构校验失败")
    if suffix == ".docx" and not zipfile.is_zipfile(io.BytesIO(data)):
        raise ValueError("DOCX 文件结构校验失败")
    return suffix


def extract_text(filename: str, data: bytes) -> dict:
    """校验并提取上传文档文本，返回来源、文本和页数。"""
    suffix = _validate(filename, data)
    pages = 0
    if suffix in {".md", ".txt"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("文本文件必须使用 UTF-8 编码") from exc
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            raise ValueError("PDF 无法读取或已损坏") from exc
        if reader.is_encrypted:
            raise ValueError("暂不支持加密 PDF")
        pages = len(reader.pages)
        if pages > MAX_PDF_PAGES:
            raise ValueError(f"PDF 页数不能超过 {MAX_PDF_PAGES} 页")
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    else:
        try:
            from docx import Document

            document = Document(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            raise ValueError("DOCX 无法读取或已损坏") from exc
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(paragraphs)

    return {"filename": Path(filename).name[:255], "text": _clean_text(text), "pages": pages}


def ingest_upload(filename: str, data: bytes) -> dict:
    """提取并索引上传文档；同名来源会被新版本替换。"""
    document = extract_text(filename, data)
    result = rag.ingest_document_text(document["text"], document["filename"])
    return {**document, **result}
