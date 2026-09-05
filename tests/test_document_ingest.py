import io
import unittest
from unittest.mock import patch

from docx import Document

from app import document_ingest, rag


class DocumentIngestTests(unittest.TestCase):
    def test_extracts_utf8_markdown(self):
        result = document_ingest.extract_text("policy.md", "# 退货政策\n\n支持七天无理由退货。".encode())

        self.assertEqual(result["filename"], "policy.md")
        self.assertIn("七天无理由退货", result["text"])
        self.assertEqual(result["pages"], 0)

    def test_extracts_docx_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph("保修期为十二个月。")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "渠道"
        table.cell(0, 1).text = "客服"
        buffer = io.BytesIO()
        document.save(buffer)

        result = document_ingest.extract_text("policy.docx", buffer.getvalue())

        self.assertIn("保修期为十二个月", result["text"])
        self.assertIn("渠道 | 客服", result["text"])

    def test_rejects_fake_pdf_and_unsupported_extension(self):
        with self.assertRaisesRegex(ValueError, "PDF 文件结构"):
            document_ingest.extract_text("policy.pdf", b"not a pdf")
        with self.assertRaisesRegex(ValueError, "仅支持"):
            document_ingest.extract_text("policy.exe", b"content")

    def test_upload_passes_extracted_text_to_rag(self):
        with patch.object(rag, "ingest_document_text", return_value={"source": "policy.md", "chunks": 2, "characters": 10}) as ingest:
            result = document_ingest.ingest_upload("policy.md", "退货政策\n\n支持七天无理由退货".encode())

        ingest.assert_called_once()
        self.assertEqual(result["source"], "policy.md")
        self.assertEqual(result["chunks"], 2)


if __name__ == "__main__":
    unittest.main()
